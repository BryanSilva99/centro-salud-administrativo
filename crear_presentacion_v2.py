from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Paleta de colores vibrante
AZUL_OSCURO = RGBColor(0, 51, 102)
AZUL_MEDIO = RGBColor(0, 102, 204)
AZUL_CLARO = RGBColor(0, 153, 255)
VERDE = RGBColor(0, 153, 76)
NARANJA = RGBColor(255, 102, 0)
ROJO = RGBColor(204, 0, 0)
MORADO = RGBColor(102, 0, 153)
ROSA = RGBColor(255, 0, 102)
AMARILLO = RGBColor(255, 204, 0)
GRIS_CLARO = RGBColor(240, 240, 240)
GRIS_MEDIO = RGBColor(200, 200, 200)
BLANCO = RGBColor(255, 255, 255)
NEGRO = RGBColor(30, 30, 30)
TURQUESA = RGBColor(0, 188, 212)

# Fondos con color
FONDO_AZUL = RGBColor(230, 242, 255)
FONDO_VERDE = RGBColor(230, 255, 230)
FONDO_NARANJA = RGBColor(255, 240, 230)
FONDO_MORADO = RGBColor(240, 230, 255)
FONDO_ROSA = RGBColor(255, 230, 240)
FONDO_AMARILLO = RGBColor(255, 255, 200)

def add_background(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_decorative_shape(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_slide(title, content, bg_color=FONDO_AZUL, title_color=AZUL_OSCURO, accent_color=AZUL_MEDIO):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, bg_color)
    
    # Forma decorativa superior
    add_decorative_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.3), accent_color)
    
    # Forma decorativa lateral
    add_decorative_shape(slide, Inches(0), Inches(0.3), Inches(0.2), Inches(7.2), accent_color)
    
    # Título con fondo
    title_bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(0.5), Inches(12.333), Inches(0.8))
    title_bg.fill.solid()
    title_bg.fill.fore_color.rgb = title_color
    title_bg.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.55), Inches(12), Inches(0.7))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = title
    title_para.font.size = Pt(32)
    title_para.font.bold = True
    title_para.font.color.rgb = BLANCO
    title_para.alignment = PP_ALIGN.LEFT
    
    # Contenido
    content_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(12), Inches(5.5))
    text_frame = content_box.text_frame
    text_frame.word_wrap = True
    
    for i, line in enumerate(content.split('\n')):
        if i == 0:
            p = text_frame.paragraphs[0]
        else:
            p = text_frame.add_paragraph()
        p.text = line
        p.font.size = Pt(20) if not line.startswith('**') else Pt(22)
        p.space_after = Pt(8)
        if line.startswith('**') and line.endswith('**'):
            p.text = line[2:-2]
            p.font.bold = True
            p.font.color.rgb = title_color
        elif line.startswith('- '):
            p.text = '  ► ' + line[2:]
            p.font.color.rgb = NEGRO
        elif line.startswith('|'):
            p.font.size = Pt(14)
            p.font.name = 'Consolas'
        else:
            p.font.color.rgb = NEGRO
    
    return slide

# Diapositiva 1 - Portada con fondo azul intenso
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, AZUL_OSCURO)

# Formas decorativas
add_decorative_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.5), AZUL_CLARO)
add_decorative_shape(slide, Inches(0), Inches(7), Inches(13.333), Inches(0.5), AZUL_CLARO)
add_decorative_shape(slide, Inches(5.5), Inches(2.5), Inches(2.333), Inches(0.1), AMARILLO)

# Título principal
title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.7), Inches(12.333), Inches(1.2))
tf = title_box.text_frame
p = tf.paragraphs[0]
p.text = "Sistema Administrativo"
p.font.size = Pt(54)
p.font.bold = True
p.font.color.rgb = BLANCO
p.alignment = PP_ALIGN.CENTER

p2 = tf.add_paragraph()
p2.text = "para Centro de Salud"
p2.font.size = Pt(40)
p2.font.bold = True
p2.font.color.rgb = AMARILLO
p2.alignment = PP_ALIGN.CENTER

# Subtítulo
sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(12.333), Inches(1))
tf2 = sub_box.text_frame
p3 = tf2.paragraphs[0]
p3.text = "Gestión integral de pacientes, citas, caja y recursos"
p3.font.size = Pt(24)
p3.font.color.rgb = RGBColor(200, 220, 255)
p3.alignment = PP_ALIGN.CENTER

# Equipo
team_box = slide.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(12.333), Inches(1))
tf3 = team_box.text_frame
p4 = tf3.paragraphs[0]
p4.text = "Equipo: Bryan, Oscar, Percy, J.J."
p4.font.size = Pt(22)
p4.font.color.rgb = BLANCO
p4.alignment = PP_ALIGN.CENTER

p5 = tf3.add_paragraph()
p5.text = "Expo Parcial - Octubre 2026"
p5.font.size = Pt(18)
p5.font.color.rgb = RGBColor(180, 200, 230)
p5.alignment = PP_ALIGN.CENTER

# Diapositiva 2 - Problema
add_slide("El Problema", """Los centros de salud manejan información dispersa:
- Papeles para citas y registros
- No hay control de pagos ni cierre de caja
- Los pacientes esperan mucho por desorganización
- No hay trazabilidad de quién hizo cada operación

**Conclusión:** Necesitamos un sistema que centralice toda la operación administrativa.""",
    bg_color=RGBColor(255, 240, 240), title_color=ROJO, accent_color=ROJO)

# Diapositiva 3 - Solución
add_slide("Nuestra Solución", """Un sistema web interno que integra:
- Gestión de pacientes y personal
- Programación de citas y atención
- Caja con cargos, pagos y cierre
- Reportes y auditoría

**Beneficios:**
- Información centralizada
- Acceso por roles (recepción, caja, administrador, profesional)
- Trazabilidad de operaciones""",
    bg_color=RGBColor(240, 255, 240), title_color=VERDE, accent_color=VERDE)

# Diapositiva 4 - Arquitectura
add_slide("Cómo Está Construido", """**Stack:**
- Frontend: React
- Backend: Java Spring Boot
- Base de datos: PostgreSQL
- Observabilidad: SigNoz

**Patrón:** Monolito modular con arquitectura hexagonal
- Dominio independiente de frameworks
- Puertos y adaptadores
- Módulos: identidad, administración, citas, caja, reportes""",
    bg_color=RGBColor(240, 240, 255), title_color=MORADO, accent_color=MORADO)

# Diapositiva 5 - Contexto
add_slide("Contexto del Sistema", """**Actores:**
- Recepción/Admisión (REC): pacientes, citas, llegadas
- Caja (CAJ): cargos, pagos, cierres
- Administrador (ADM): configuración, usuarios, reportes
- Profesional de salud (PRO): agenda, atención

**El paciente NO es usuario** del sistema.""",
    bg_color=RGBColor(255, 250, 240), title_color=NARANJA, accent_color=NARANJA)

# Diapositiva 6 - Módulos
add_slide("Módulos Implementados", """| Módulo | Responsable | Funcionalidad |
|--------|-------------|---------------|
| Identidad y acceso | Bryan | Login, roles, permisos, auditoría |
| Administración | Oscar | Pacientes, profesionales, servicios, consultorios, horarios |
| Citas y atención | Percy | Reservas, reprogramaciones, llegadas, atención |
| Caja | J.J. | Cargos, pagos, anulaciones, cierre, reportes |""",
    bg_color=RGBColor(245, 255, 250), title_color=TURQUESA, accent_color=TURQUESA)

# Diapositiva 7 - Requerimientos
add_slide("Lo Que el Negocio Necesita (P8)", """1. **Control de acceso** — Cada rol solo ve lo que necesita
2. **Citas sin conflictos** — Un profesional no puede tener dos citas a la vez
3. **Caja con conciliación** — Al cerrar, el efectivo debe cuadrar
4. **Sin duplicados** — Un paciente no se registra dos veces
5. **Trazabilidad** — Todo queda registrado con quién, cuándo y qué""",
    bg_color=RGBColor(250, 245, 255), title_color=ROSA, accent_color=ROSA)

# Diapositiva 8 - RF
add_slide("RF Implementados", """- **RF-19 a RF-24:** Caja completa (cargos, pagos, saldos, anulaciones, cierre)
- **RF-27:** Reporte económico
- **RF-04 a RF-10:** Pacientes, profesionales, servicios, consultorios, disponibilidad
- **RF-12 a RF-18:** Citas, reprogramaciones, llegadas, atención""",
    bg_color=RGBColor(240, 255, 255), title_color=AZUL_MEDIO, accent_color=AZUL_MEDIO)

# Diapositiva 9 - Reglas
add_slide("Reglas Que el Sistema Aplica", """- **RN-07:** No se pueden reservar citas que se crucen (mismo paciente, profesional o consultorio)
- **RN-13:** Los pagos no pueden exceder el saldo pendiente
- **RN-14:** El solicitante de anulación no puede aprobarla
- **RN-15:** Una sola caja abierta por cajero; el cierre es inmutable""",
    bg_color=RGBColor(255, 245, 250), title_color=RGBColor(180, 80, 0), accent_color=RGBColor(180, 80, 0))

# Diapositiva 10 - Tendencias
add_slide("6 Tendencias Tecnológicas", """| # | Tendencia | Categoría |
|---|-----------|-----------|
| 1 | Zero Trust Architecture | Techniques |
| 2 | Mutation Testing | Techniques |
| 3 | OpenAI Codex | Tools |
| 4 | Axe-core | Tools |
| 5 | SigNoz | Platforms |
| 6 | React JS | Languages and Frameworks |

**Distribución:** 2 + 2 + 1 + 1""",
    bg_color=RGBColor(250, 255, 250), title_color=VERDE, accent_color=VERDE)

# Diapositiva 11 - Zero Trust
add_slide("Zero Trust Architecture", """**Qué es:** Verificar cada operación, sin confiar en la red ni en el navegador.

**Cómo lo implementamos:**
- Autenticación por sesión con cookies seguras (HttpOnly, SameSite, Secure)
- Autorización por rol Y permiso en cada endpoint
- CSRF tokens en operaciones de escritura
- Límite de intentos de login (5 en 15 minutos)
- El servidor valida todo, no confía en el cliente""",
    bg_color=RGBColor(240, 250, 255), title_color=AZUL_OSCURO, accent_color=AZUL_OSCURO)

# Diapositiva 12 - Mutation Testing
add_slide("Mutation Testing", """**Qué es:** Verificar que las pruebas detectan errores. Se introduce código con errores (mutantes) y se comprueba que las pruebas fallan.

**Cómo lo aplicamos:**
- Pruebas de reglas críticas: superposición de citas, saldos, anulaciones
- Si una mutación sobrevive, la prueba es insuficiente
- Herramienta: PIT (pendiente de ejecutar)""",
    bg_color=RGBColor(255, 250, 255), title_color=MORADO, accent_color=MORADO)

# Diapositiva 13 - Herramientas
add_slide("Herramientas de Desarrollo", """**OpenAI Codex:**
- Asistió en la generación de código del módulo de caja
- Se usó con revisión humana
- Evidencia: commits y documentación en docs/

**Axe-core:**
- Detecta problemas de accesibilidad en React
- Se integra en el build del frontend
- Verifica etiquetas, foco visible, contraste""",
    bg_color=RGBColor(255, 240, 250), title_color=ROSA, accent_color=ROSA)

# Diapositiva 14 - Plataforma y Framework
add_slide("Plataforma y Framework", """**SigNoz:**
- Observabilidad: métricas, logs y trazas
- Correlación de operaciones de reserva y pago
- No bloquea el negocio si se cae

**React JS:**
- Interfaz interna para los 4 roles
- Componentes reutilizables
- Navegación por permisos""",
    bg_color=RGBColor(240, 255, 250), title_color=TURQUESA, accent_color=TURQUESA)

# Diapositiva 15 - Matriz
add_slide("Tendencias ↔ Requerimientos", """| Tendencia | RF-01 a 03 | RF-04 a 10 | RF-12 a 18 | RF-19 a 24 | RF-27 | RNF-01 |
|-----------|------------|------------|------------|------------|-------|--------|
| Zero Trust | ✓ | ✓ | ✓ | ✓ | ✓ | ✓ |
| Mutation Testing | — | ✓ | ✓ | ✓ | ✓ | ✓ |
| OpenAI Codex | — | — | ✓ | ✓ | ✓ | — |
| Axe-core | — | — | — | — | — | — |
| SigNoz | — | — | ✓ | ✓ | ✓ | ✓ |
| React JS | ✓ | ✓ | ✓ | ✓ | ✓ | ✓ |""",
    bg_color=RGBColor(245, 245, 255), title_color=AZUL_MEDIO, accent_color=AZUL_MEDIO)

# Diapositiva 16 - ADD Contexto
add_slide("ADD — Contexto del Sistema", """**Sistema:** Administración integral de un centro de salud

**Contenedores:**
1. Aplicación web (React)
2. API Spring Boot (monolito modular)
3. PostgreSQL
4. SigNoz (observabilidad)

**Restricciones:**
- Sin historia clínica
- Sin portal de paciente
- Sin facturación electrónica
- Una sola sede""",
    bg_color=RGBColor(255, 250, 245), title_color=NARANJA, accent_color=NARANJA)

# Diapositiva 17 - ADD Arquitectura
add_slide("ADD — Arquitectura del Software", """**Patrón:** Monolito modular con arquitectura hexagonal

HTTP (Controladores) → Casos de Uso → Dominio
                                    ↓
                              Puertos
                                    ↓
                         Adaptadores (PostgreSQL, SigNoz)

**Decisiones:**
- DA-01: Monolito modular (transacciones que abarcan módulos)
- DA-02: Arquitectura hexagonal (dominio independiente)
- DA-04: Una transacción por operación de negocio
- DA-05: Movimientos compensatorios, nunca borrado""",
    bg_color=RGBColor(245, 250, 255), title_color=MORADO, accent_color=MORADO)

# Diapositiva 18 - ADD Vistas
add_slide("ADD — Vistas del Sistema", """**Vista Lógica:**
- Módulo de Identidad
- Módulo de Administración
- Módulo de Citas
- Módulo de Caja
- Módulo de Reportes

**Vista de Despliegue:**
- Servidor web → API → PostgreSQL
- API → SigNoz (telemetría asíncrona)
- Sin estado local en memoria""",
    bg_color=RGBColor(240, 250, 250), title_color=TURQUESA, accent_color=TURQUESA)

# Diapositiva 19 - Ejecución
add_slide("Demostración en Vivo", """**Flujos a mostrar:**
1. Login con rol de caja
2. Abrir sesión de caja
3. Registrar cargo por servicio
4. Registrar pago parcial y total
5. Consultar saldo
6. Cerrar caja con conciliación
7. Ver reporte económico

**Comando para levantar:**
docker compose up --build""",
    bg_color=RGBColor(255, 245, 240), title_color=ROJO, accent_color=ROJO)

# Diapositiva 20 - Conclusiones
add_slide("Conclusiones", """**Zero Trust:**
- Garantiza que cada operación está autorizada
- Protege datos sensibles del centro

**Mutation Testing:**
- Asegura que las pruebas realmente verifican las reglas críticas

**OpenAI Codex:**
- Acelera el desarrollo con revisión humana

**Axe-core:**
- Garantiza accesibilidad para todo el personal

**SigNoz:**
- Permite diagnosticar problemas sin afectar operaciones

**React JS:**
- Interfaz moderna y mantenible""",
    bg_color=RGBColor(250, 250, 255), title_color=VERDE, accent_color=VERDE)

# Diapositiva 21 - Cierre con estilo
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, AZUL_OSCURO)

# Formas decorativas
add_decorative_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.5), AZUL_CLARO)
add_decorative_shape(slide, Inches(0), Inches(7), Inches(13.333), Inches(0.5), AZUL_CLARO)
add_decorative_shape(slide, Inches(5.5), Inches(3), Inches(2.333), Inches(0.1), AMARILLO)

# Gracias
title_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.2), Inches(12.333), Inches(1))
tf = title_box.text_frame
p = tf.paragraphs[0]
p.text = "¡Gracias!"
p.font.size = Pt(60)
p.font.bold = True
p.font.color.rgb = BLANCO
p.alignment = PP_ALIGN.CENTER

# Equipo
sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.5), Inches(12.333), Inches(1.5))
tf2 = sub_box.text_frame
p2 = tf2.paragraphs[0]
p2.text = "Preguntas y Respuestas"
p2.font.size = Pt(28)
p2.font.color.rgb = RGBColor(200, 220, 255)
p2.alignment = PP_ALIGN.CENTER

p3 = tf2.add_paragraph()
p3.text = "Equipo: Bryan, Oscar, Percy, J.J."
p3.font.size = Pt(24)
p3.font.color.rgb = BLANCO
p3.alignment = PP_ALIGN.CENTER

p4 = tf2.add_paragraph()
p4.text = "Centro de Salud - Expo Parcial 2026"
p4.font.size = Pt(18)
p4.font.color.rgb = RGBColor(180, 200, 230)
p4.alignment = PP_ALIGN.CENTER

prs.save('C:\\Users\\Juanjo\\centro-salud-jj\\Presentacion_Expo_Centro_Salud_v2.pptx')
print("Presentación v2 creada: Presentacion_Expo_Centro_Salud_v2.pptx")
