from pptx import Presentation
import sys

sys.stdout.reconfigure(encoding='utf-8')
prs = Presentation(r'C:\Users\Juanjo\Downloads\Presentacion_Expo_Centro_Salud_v3.pptx')

for i, slide in enumerate(prs.slides):
    print(f'=== Diapositiva {i+1} ===')
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            print(shape.text_frame.text)
    print()
