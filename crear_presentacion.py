from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_slide(title, content, is_title=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Título
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.333), Inches(1))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = title
    title_para.font.size = Pt(36) if is_title else Pt(28)
    title_para.font.bold = True
    title_para.font.color.rgb = RGBColor(0, 102, 204) if not is_title else RGBColor(255, 255, 255)
    if is_title:
        title_para.alignment = PP_ALIGN.CENTER
    
    # Contenido
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12.333), Inches(5.5))
    text_frame = content_box.text_frame
    text_frame.word_wrap = True
    
    for i, line in enumerate(content.split('\n')):
        if i == 0:
            p = text_frame.paragraphs[0]
        else:
            p = text_frame.add_paragraph()
        p.text = line
        p.font.size = Pt(18) if not is_title else Pt(24)
        p.space_after = Pt(6)
        if line.startswith('**') and line.endswith('**'):
            p.text = line[2:-2]
            p.font.bold = True
        elif line.startswith('- '):
            p.text = '  • ' + line[2:]
    
    return slide

# Diapositiva 1 - Portada
slide = prs.slides.add_slide(prs.slide_layouts[6])
title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(12.333), Inches(1.5))
tf = title_box.text_frame
p = tf.paragraphs[0]
p.text = "Sistema Administrativo para Centro de Salud"
p.font.size = Pt(44)
p.font.bold = True
p.font.color.rgb = RGBColor(0, 102, 204)
p.alignment = PP_ALIGN.CENTER

sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.5), Inches(12.333), Inches(2))
tf2 = sub_box.text_frame
p2 = tf2.paragraphs[0]
p2.text = "Gestión integral de pacientes, citas, caja y recursos"
p2.font.size = Pt(24)
p2.alignment = PP_ALIGN.CENTER
p3 = tf2.add_paragraph()
p3.text = "Equipo: Bryan, Oscar, Percy, J.J."
p3.font.size = Pt(20)
p3.alignment = PP_ALIGN.CENTER
p4 = tf2.add_paragraph()
p4.text = "Octubre 2026"
p4.font.size = Pt(18)
p4.alignment = PP_ALIGN.CENTER

# Diapositiva 2 - Problema
add_slide("El Problema", """Los centros de salud manejan información dispersa:
- Papeles para citas y registros
- No hay control de pagos ni cierre de caja
- Los pacientes esperan mucho por desorganización
- No hay trazabilidad de quién hizo cada operación

**Conclusión:** Necesitamos un sistema que centralice toda la operación administrativa.""")

# Diapositiva 3 - Solución
add_slide("Nuestra Solución", """Un sistema web interno que integra:
- Gestión de pacientes y personal
- Programación de citas y atención
- Caja con cargos, pagos y cierre
- Reportes y auditoría

**Beneficios:**
- Información centralizada
- Acceso por roles (recepción, caja, administrador, profesional)
- Trazabilidad de operaciones""")

# Diapositiva 4 - Arquitectura
add_slide("Cómo Está Construido", """**Stack:**
- Frontend: React
- Backend: Java Spring Boot
- Base de datos: PostgreSQL
- Observabilidad: SigNoz

**Patrón:** Monolito modular con arquitectura hexagonal
- Dominio independiente de frameworks
- Puertos y adaptadores
- Módulos: identidad, administración, citas, caja, reportes""")

# Diapositiva 5 - Contexto del Sistema
add_slide("Contexto del Sistema", """**Actores:**
- Recepción/Admisión (REC): pacientes, citas, llegadas
- Caja (CAJ): cargos, pagos, cierres
- Administrador (ADM): configuración, usuarios, reportes
- Profesional de salud (PRO): agenda, atención

**El paciente NO es usuario** del sistema.""")

# Diapositiva 6 - Módulos
add_slide("Módulos Implementados", """| Módulo | Responsable | Funcionalidad |
|--------|-------------|---------------|
| Identidad y acceso | Bryan | Login, roles, permisos, auditoría |
| Administración | Oscar | Pacientes, profesionales, servicios, consultorios, horarios |
| Citas y atención | Percy | Reservas, reprogramaciones, llegadas, atención |
| Caja | J.J. | Cargos, pagos, anulaciones, cierre, reportes |""")

# Diapositiva 7 - Requerimientos del Negocio
add_slide("Lo Que el Negocio Necesita (P8)", """1. **Control de acceso** — Cada rol solo ve lo que necesita
2. **Citas sin conflictos** — Un profesional no puede tener dos citas a la vez
3. **Caja con conciliación** — Al cerrar, el efectivo debe cuadrar
4. **Sin duplicados** — Un paciente no se registra dos veces
5. **Trazabilidad** — Todo queda registrado con quién, cuándo y qué""")

# Diapositiva 8 - RF Clave
add_slide("RF Implementados", """- **RF-19 a RF-24:** Caja completa (cargos, pagos, saldos, anulaciones, cierre)
- **RF-27:** Reporte económico
- **RF-04 a RF-10:** Pacientes, profesionales, servicios, consultorios, disponibilidad
- **RF-12 a RF-18:** Citas, reprogramaciones, llegadas, atención""")

# Diapositiva 9 - Reglas de Negocio
add_slide("Reglas Que el Sistema Aplica", """- **RN-07:** No se pueden reservar citas que se crucen (mismo paciente, profesional o consultorio)
- **RN-13:** Los pagos no pueden exceder el saldo pendiente
- **RN-14:** El solicitante de anulación no puede aprobarla
- **RN-15:** Una sola caja abierta por cajero; el cierre es inmutable""")

# Diapositiva 10 - Tendencias
add_slide("6 Tendencias Tecnológicas", """| # | Tendencia | Categoría |
|---|-----------|-----------|
| 1 | Zero Trust Architecture | Techniques |
| 2 | Mutation Testing | Techniques |
| 3 | OpenAI Codex | Tools |
| 4 | Axe-core | Tools |
| 5 | SigNoz | Platforms |
| 6 | React JS | Languages and Frameworks |

**Distribución:** 2 + 2 + 1 + 1""")

# Diapositiva 11 - Zero Trust
add_slide("Zero Trust Architecture", """**Qué es:** Verificar cada operación, sin confiar en la red ni en el navegador.

**Cómo lo implementamos:**
- Autenticación por sesión con cookies seguras (HttpOnly, SameSite, Secure)
- Autorización por rol Y permiso en cada endpoint
- CSRF tokens en operaciones de escritura
- Límite de intentos de login (5 en 15 minutos)
- El servidor valida todo, no confía en el cliente""")

# Diapositiva 12 - Mutation Testing
add_slide("Mutation Testing", """**Qué es:** Verificar que las pruebas detectan errores. Se introduce código con errores (mutantes) y se comprueba que las pruebas fallan.

**Cómo lo aplicamos:**
- Pruebas de reglas críticas: superposición de citas, saldos, anulaciones
- Si una mutación sobrevive, la prueba es insuficiente
- Herramienta: PIT (pendiente de ejecutar)""")

# Diapositiva 13 - OpenAI Codex y Axe-core
add_slide("Herramientas de Desarrollo", """**OpenAI Codex:**
- Asistió en la generación de código del módulo de caja
- Se usó con revisión humana
- Evidencia: commits y documentación en docs/

**Axe-core:**
- Detecta problemas de accesibilidad en React
- Se integra en el build del frontend
- Verifica etiquetas, foco visible, contraste""")

# Diapositiva 14 - SigNoz y React
add_slide("Plataforma y Framework", """**SigNoz:**
- Observabilidad: métricas, logs y trazas
- Correlación de operaciones de reserva y pago
- No bloquea el negocio si se cae

**React JS:**
- Interfaz interna para los 4 roles
- Componentes reutilizables
- Navegación por permisos""")

# Diapositiva 15 - Matriz
add_slide("Tendencias ↔ Requerimientos", """| Tendencia | RF-01 a 03 | RF-04 a 10 | RF-12 a 18 | RF-19 a 24 | RF-27 | RNF-01 |
|-----------|------------|------------|------------|------------|-------|--------|
| Zero Trust | ✓ | ✓ | ✓ | ✓ | ✓ | ✓ |
| Mutation Testing | — | ✓ | ✓ | ✓ | ✓ | ✓ |
| OpenAI Codex | — | — | ✓ | ✓ | ✓ | — |
| Axe-core | — | — | — | — | — | — |
| SigNoz | — | — | ✓ | ✓ | ✓ | ✓ |
| React JS | ✓ | ✓ | ✓ | ✓ | ✓ | ✓ |""")

# Diapositiva 16 - ADD Contexto
add_slide("ADD — Contexto del Sistema", """**Sistema:** Administración integral de un centro de salud

**Contenedores:**
1. Aplicación web (React)
2. API Spring Boot (monolito modular)
3. PostgreSQL
4. SigNoz (observabilidad)

**Restricciones:**
- Sin historia clínica
- Sin portal de paciente
- Sin facturación electrónica
- Una sola sede""")

# Diapositiva 17 - ADD Arquitectura
add_slide("ADD — Arquitectura del Software", """**Patrón:** Monolito modular con arquitectura hexagonal

HTTP (Controladores) → Casos de Uso → Dominio
                                    ↓
                              Puertos
                                    ↓
                         Adaptadores (PostgreSQL, SigNoz)

**Decisiones:**
- DA-01: Monolito modular (transacciones que abarcan módulos)
- DA-02: Arquitectura hexagonal (dominio independiente)
- DA-04: Una transacción por operación de negocio
- DA-05: Movimientos compensatorios, nunca borrado""")

# Diapositiva 18 - ADD Vistas
add_slide("ADD — Vistas del Sistema", """**Vista Lógica:**
- Módulo de Identidad
- Módulo de Administración
- Módulo de Citas
- Módulo de Caja
- Módulo de Reportes

**Vista de Despliegue:**
- Servidor web → API → PostgreSQL
- API → SigNoz (telemetría asíncrona)
- Sin estado local en memoria""")

# Diapositiva 19 - Ejecución
add_slide("Demostración en Vivo", """**Flujos a mostrar:**
1. Login con rol de caja
2. Abrir sesión de caja
3. Registrar cargo por servicio
4. Registrar pago parcial y total
5. Consultar saldo
6. Cerrar caja con conciliación
7. Ver reporte económico

**Comando para levantar:**
docker compose up --build""")

# Diapositiva 20 - Conclusiones
add_slide("Conclusiones", """**Zero Trust:**
- Garantiza que cada operación está autorizada
- Protege datos sensibles del centro

**Mutation Testing:**
- Asegura que las pruebas realmente verifican las reglas críticas

**OpenAI Codex:**
- Acelera el desarrollo con revisión humana

**Axe-core:**
- Garantiza accesibilidad para todo el personal

**SigNoz:**
- Permite diagnosticar problemas sin afectar operaciones

**React JS:**
- Interfaz moderna y mantenible""")

# Diapositiva 21 - Cierre
slide = prs.slides.add_slide(prs.slide_layouts[6])
title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(1))
tf = title_box.text_frame
p = tf.paragraphs[0]
p.text = "Gracias"
p.font.size = Pt(48)
p.font.bold = True
p.font.color.rgb = RGBColor(0, 102, 204)
p.alignment = PP_ALIGN.CENTER

sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(4), Inches(12.333), Inches(2))
tf2 = sub_box.text_frame
p2 = tf2.paragraphs[0]
p2.text = "Preguntas y Respuestas"
p2.font.size = Pt(28)
p2.alignment = PP_ALIGN.CENTER
p3 = tf2.add_paragraph()
p3.text = "Equipo: Bryan, Oscar, Percy, J.J."
p3.font.size = Pt(20)
p3.alignment = PP_ALIGN.CENTER

prs.save('C:\\Users\\Juanjo\\centro-salud-jj\\Presentacion_Expo_Centro_Salud.pptx')
print("Presentación creada exitosamente: Presentacion_Expo_Centro_Salud.pptx")
